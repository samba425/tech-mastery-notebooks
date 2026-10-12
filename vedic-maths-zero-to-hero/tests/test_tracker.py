"""The Excel tracker builds, has the expected sheets, and its dropdowns stay within Excel's limits."""
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import generate_tracker  # noqa: E402


def test_tracker(tmp_path):
    out = tmp_path / "tracker.xlsx"
    generate_tracker.build(out)
    wb = load_workbook(out)
    assert wb.sheetnames[:6] == ["Daily Log", "Practice Sets", "Mock Tests", "Mistake Log", "Topic Mastery", "90-Day Plan"]
    sets = wb["Practice Sets"]
    assert [sets.cell(row=r, column=1).value for r in (4, 73)] == [1, 70]
    assert wb["90-Day Plan"].cell(row=93, column=1).value == 90
    for ws in wb:
        for dv in ws.data_validations.dataValidation:
            assert len(dv.formula1) <= 255
            assert dv.formula1.startswith("Lists!")
