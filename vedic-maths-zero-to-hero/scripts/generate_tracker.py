#!/usr/bin/env python3
"""Build output/Vedic_Maths_Progress_Tracker.xlsx (practice log, practice sets, mocks, mistakes, mastery, 90-day plan)."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from booklib.generated import DIFF, SECONDS, SET_CATEGORIES  # noqa: E402

OUT = Path(__file__).resolve().parents[1] / "output" / "Vedic_Maths_Progress_Tracker.xlsx"
ROWS = 200

HEAD = PatternFill("solid", fgColor="0F766E")
ZEBRA = PatternFill("solid", fgColor="F1F5F9")
TITLE = Font(name="Calibri", size=16, bold=True, color="1E40AF")
WHITE = Font(bold=True, color="FFFFFF")
THIN = Border(bottom=Side(style="thin", color="CBD5E1"))
GOOD = PatternFill("solid", fgColor="DCFCE7")
BAD = PatternFill("solid", fgColor="FEE2E2")

TOPICS = ["Complements and add/sub", "×11, ×5, ×25, ×9s", "Nikhilam", "Urdhva-Tiryagbhyam", "Squares and square roots",
          "Cubes and cube roots", "Division methods", "Fractions and percentages", "Algebra identities and factorising",
          "Simultaneous and quadratic equations", "Ratio, averages, mixtures", "Profit, loss, discount", "Simple and compound interest",
          "Time and work, pipes", "Speed, trains, boats", "Number system and remainders", "Counting and probability",
          "Mensuration", "Calendar", "Simplification and approximation", "Data interpretation"]
REASONS = ["Concept gap", "Careless", "Too slow", "Misread question", "Wrong shortcut", "Calculation slip", "Guessed"]


def sheet(wb, name, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name
    ws["A1"] = title
    ws["A1"].font = TITLE
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.fill, cell.font = HEAD, WHITE
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "A4"
    return ws


def zebra(ws, ncols, start, end):
    for r in range(start, end + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = THIN
            if r % 2 == 0:
                cell.fill = ZEBRA


def pct_rules(ws, rng, good=0.9, bad=0.7):
    ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=[str(good)], fill=GOOD))
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[str(bad)], fill=BAD))


LISTS: dict[str, str] = {}


def dropdown(ws, options, rng):
    """Excel caps inline list formulas at 255 characters, so options live on a hidden 'Lists' sheet."""
    key = "|".join(options)
    if key not in LISTS:
        lists = ws.parent["Lists"] if "Lists" in ws.parent.sheetnames else ws.parent.create_sheet("Lists")
        lists.sheet_state = "hidden"
        col = get_column_letter(len(LISTS) + 1)
        for i, opt in enumerate(options, start=1):
            lists[f"{col}{i}"] = opt
        LISTS[key] = f"Lists!${col}$1:${col}${len(options)}"
    dv = DataValidation(type="list", formula1=LISTS[key], allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def build(out: Path = OUT):
    LISTS.clear()
    wb = Workbook()

    ws = sheet(wb, "Daily Log", "Daily practice log", ["Date", "Topic", "Questions", "Correct", "Accuracy", "Minutes", "Sec / question", "Notes"],
               [12, 34, 11, 10, 11, 10, 14, 40], first=True)
    for r in range(4, ROWS + 4):
        ws[f"E{r}"] = f'=IF(C{r}>0,D{r}/C{r},"")'
        ws[f"G{r}"] = f'=IF(C{r}>0,F{r}*60/C{r},"")'
        ws[f"E{r}"].number_format = "0%"
        ws[f"G{r}"].number_format = "0"
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
    dropdown(ws, [t.replace(",", ";") for t in TOPICS], f"B4:B{ROWS + 3}")
    pct_rules(ws, f"E4:E{ROWS + 3}")
    zebra(ws, 8, 4, ROWS + 3)

    ws = sheet(wb, "Practice Sets", "Part XIII practice sets (8 questions each)",
               ["Set", "Category", "Difficulty", "Suggested min", "Score / 8", "Minutes taken", "Accuracy", "Within time?", "Date"],
               [7, 34, 14, 14, 11, 14, 11, 13, 12])
    r = 4
    n = 0
    for cat, _ in SET_CATEGORIES:
        for i in range(10):
            n += 1
            diff = DIFF[i]
            ws.append([n, cat, diff, max(4, round(8 * SECONDS[diff] / 60))])
            ws[f"G{r}"] = f'=IF(E{r}<>"",E{r}/8,"")'
            ws[f"G{r}"].number_format = "0%"
            ws[f"H{r}"] = f'=IF(F{r}<>"",IF(F{r}<=D{r},"Yes","No"),"")'
            r += 1
    pct_rules(ws, f"G4:G{r - 1}")
    zebra(ws, 9, 4, r - 1)

    ws = sheet(wb, "Mock Tests", "Mock test record",
               ["Date", "Test name", "Questions", "Attempted", "Correct", "Wrong", "Marks / correct", "Negative / wrong", "Score", "Accuracy",
                "Minutes", "Weakest topic"],
               [12, 26, 11, 11, 10, 10, 14, 15, 10, 11, 10, 30])
    for r in range(4, 104):
        ws[f"F{r}"] = f'=IF(D{r}<>"",D{r}-E{r},"")'
        ws[f"I{r}"] = f'=IF(D{r}<>"",E{r}*G{r}-F{r}*H{r},"")'
        ws[f"J{r}"] = f'=IF(D{r}>0,E{r}/D{r},"")'
        ws[f"J{r}"].number_format = "0%"
        ws[f"A{r}"].number_format = "yyyy-mm-dd"
    dropdown(ws, [t.replace(",", ";") for t in TOPICS], "L4:L103")
    pct_rules(ws, "J4:J103")
    zebra(ws, 12, 4, 103)
    ws["N3"], ws["N4"], ws["N5"], ws["N6"] = "Summary", "Mocks taken", "Average score", "Average accuracy"
    ws["N3"].font = Font(bold=True)
    ws["O4"] = '=COUNT(I4:I103)'
    ws["O5"] = '=IFERROR(AVERAGE(I4:I103),"")'
    ws["O6"] = '=IFERROR(AVERAGE(J4:J103),"")'
    ws["O6"].number_format = "0%"
    ws.column_dimensions["N"].width = 18

    ws = sheet(wb, "Mistake Log", "Mistake log: redo every row a week later",
               ["Date", "Source (book / set / mock)", "Question no.", "Topic", "Reason", "What I will do differently", "Redone?"],
               [12, 24, 12, 34, 18, 44, 10])
    dropdown(ws, [t.replace(",", ";") for t in TOPICS], f"D4:D{ROWS + 3}")
    dropdown(ws, REASONS, f"E4:E{ROWS + 3}")
    dropdown(ws, ["Yes", "No"], f"G4:G{ROWS + 3}")
    zebra(ws, 7, 4, ROWS + 3)

    ws = sheet(wb, "Topic Mastery", "Topic mastery: accuracy is calculated from the Daily Log",
               ["Topic", "Questions logged", "Correct", "Accuracy", "Status", "Date mastered"], [38, 16, 12, 12, 16, 15])
    for i, t in enumerate(TOPICS):
        r = 4 + i
        key = t.replace(",", ";")
        ws[f"A{r}"] = key
        ws[f"B{r}"] = f"=SUMIF('Daily Log'!B:B,A{r},'Daily Log'!C:C)"
        ws[f"C{r}"] = f"=SUMIF('Daily Log'!B:B,A{r},'Daily Log'!D:D)"
        ws[f"D{r}"] = f'=IF(B{r}>0,C{r}/B{r},"")'
        ws[f"D{r}"].number_format = "0%"
        ws[f"E{r}"] = f'=IF(B{r}<20,"Need more data",IF(D{r}>=0.9,"Mastered",IF(D{r}>=0.75,"Getting there","Revise")))'
    end = 3 + len(TOPICS)
    ws.conditional_formatting.add(f"D4:D{end}", ColorScaleRule(start_type="num", start_value=0.5, start_color="FCA5A5",
                                                                end_type="num", end_value=1, end_color="86EFAC"))
    zebra(ws, 6, 4, end)

    ws = sheet(wb, "90-Day Plan", "90-day plan for government exams (see Part XIV)", ["Day", "Phase", "Focus", "Done?", "Notes"], [7, 14, 60, 9, 40])
    phases = [(1, 15, "Foundation", "Parts I–III; tables to 20; squares to 30; cubes to 15"),
              (16, 35, "Speed", "Parts IV–VI; one timed test from Part XII daily"),
              (36, 65, "Aptitude", "Part VIII (two chapters a week) plus one practice set daily"),
              (66, 80, "Exam skills", "Part XI; DI and approximation daily; sectional mocks"),
              (81, 90, "Revision", "Part XV formula sheets; redo the Mistake Log; full mocks")]
    for a, b, ph, focus in phases:
        for d in range(a, b + 1):
            ws.append([d, ph, focus + (" · weekly revision" if d % 7 == 0 else "")])
    dropdown(ws, ["Yes", "No"], "D4:D93")
    zebra(ws, 5, 4, 93)

    wb.move_sheet("Lists", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("Lists"))
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"tracker  {out.name}  sheets={len(wb.sheetnames)}")


if __name__ == "__main__":
    build()
